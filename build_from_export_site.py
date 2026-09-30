#!/usr/bin/env python3
"""Generate the local solar dashboard data from KWH.xlsx.

V2 goals:
- keep the workbook as the single source of truth;
- derive the consumption matrix from daily sheets (not from the Graf helper sheet);
- expose technical performance of the two PV phases;
- separate annual energy balance, self-consumption and self-sufficiency;
- seasonally project the current year from the first complete reference year;
- surface data-quality/estimation notes instead of hiding approximations.

The dashboard stays static on purpose: this script writes site/data.js so the
site can be opened directly from disk without a web server.
"""
from __future__ import annotations

import calendar
import json
import math
import sys
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    print("Erreur: le module Python 'openpyxl' est introuvable.")
    print("Installez-le avec: python -m pip install openpyxl")
    raise SystemExit(1) from exc

MONTH_NAMES = [
    "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
    "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre",
]
FALLBACK_FIRST_FULL_SOLAR_YEAR = 2025


def normalize_label(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower()
    for token in ["€", "/", "\\", "+", "-", "_", ".", ",", "(", ")", "[", "]"]:
        text = text.replace(token, " ")
    return " ".join(text.split())


def to_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return default
        return float(value)
    text = str(value).strip().replace("\u202f", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def to_optional_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    number = to_float(value, math.nan)
    return None if math.isnan(number) or math.isinf(number) else number


def to_int(value: Any, default: int = 0) -> int:
    try:
        return int(round(to_float(value, float(default))))
    except (ValueError, OverflowError):
        return default


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 1 <= float(value) <= 100000:
        # Excel 1900 date system (with the conventional 1899-12-30 origin).
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    if isinstance(value, str) and value.strip():
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def date_to_json(value: date | None) -> str | None:
    return value.isoformat() if value else None


def safe_round(value: float | None, digits: int = 2) -> float | None:
    if value is None or math.isnan(value) or math.isinf(value):
        return None
    return round(float(value), digits)


def days_in_year(year: int) -> int:
    return 366 if calendar.isleap(year) else 365


def add_year_fraction(start: date, years: float) -> date | None:
    if years <= 0 or math.isinf(years) or math.isnan(years):
        return None
    return start + timedelta(days=round(years * 365.25))


def weighted_date(phases: list[dict[str, Any]]) -> date | None:
    valid = [(as_date(p.get("date")), to_float(p.get("cost_eur"))) for p in phases]
    valid = [(d, c) for d, c in valid if d and c > 0]
    total = sum(c for _, c in valid)
    if not valid or total <= 0:
        return None
    ordinal = round(sum(d.toordinal() * c for d, c in valid) / total)
    return date.fromordinal(ordinal)


def read_pairs(workbook: Any, sheet_name: str, start_row: int = 1) -> dict[str, Any]:
    if sheet_name not in workbook.sheetnames:
        return {}
    ws = workbook[sheet_name]
    pairs: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=start_row, min_col=1, max_col=2):
        key = row[0].value
        if key is None or str(key).strip() == "":
            continue
        pairs[str(key).strip()] = row[1].value
    return pairs


def scan_kit_years(workbook: Any) -> dict[int, dict[str, float]]:
    if "Kit_Solaire" not in workbook.sheetnames:
        return {}
    ws = workbook["Kit_Solaire"]
    years: dict[int, dict[str, float]] = {}
    for row in range(1, ws.max_row + 1):
        year = to_int(ws.cell(row=row, column=2).value, -1)
        if 2000 <= year <= 2100:
            years[year] = {
                "theoretical_savings_eur": to_float(ws.cell(row=row, column=3).value),
                "workbook_savings_eur": to_float(ws.cell(row=row, column=4).value),
                "production_kwh": to_float(ws.cell(row=row, column=5).value),
                "remaining_workbook_eur": to_float(ws.cell(row=row, column=6).value),
                "injection_kwh": to_float(ws.cell(row=row, column=9).value),
                "avg_saved_kwh_price": to_float(ws.cell(row=row, column=10).value),
                "lost_injection_value_eur": to_float(ws.cell(row=row, column=11).value),
                "resale_eur": to_float(ws.cell(row=row, column=12).value),
                "theoretical_resale_eur": to_float(ws.cell(row=row, column=14).value),
                "factured_resale_kwh": to_float(ws.cell(row=row, column=15).value),
            }
    return years


def detect_header(ws: Any) -> tuple[int, dict[str, int]]:
    for row_index in range(1, min(ws.max_row, 12) + 1):
        labels = {
            normalize_label(ws.cell(row=row_index, column=col).value): col
            for col in range(1, ws.max_column + 1)
            if ws.cell(row=row_index, column=col).value is not None
        }
        if "date" in labels and any(label == "kwh" for label in labels):
            return row_index, labels
    return 1, {}


def pick_column(labels: dict[str, int], *predicates: str) -> int | None:
    for label, col in labels.items():
        if all(part in label for part in predicates):
            return col
    return None


def read_daily_sheet(workbook: Any, year: int, main_service_date: date | None) -> list[dict[str, Any]]:
    name = str(year)
    if name not in workbook.sheetnames:
        return []
    ws = workbook[name]
    header_row, labels = detect_header(ws)
    date_col = labels.get("date")
    kwh_col = labels.get("kwh") or pick_column(labels, "kwh")
    conso_col = pick_column(labels, "conso")
    subscription_col = pick_column(labels, "abonnement")
    total_col = pick_column(labels, "total", "j") or pick_column(labels, "total")
    prod_total_col = pick_column(labels, "prod solaire total")
    prod_plug_col = pick_column(labels, "prod solaire", "cab")
    prod_main_col = pick_column(labels, "prod solaire", "maison")
    prod_single_col = pick_column(labels, "production solaire") or pick_column(labels, "prod solaire", "kwh")
    if not date_col or not kwh_col:
        return []

    rows: list[dict[str, Any]] = []
    for row_index in range(header_row + 1, ws.max_row + 1):
        row_date = as_date(ws.cell(row=row_index, column=date_col).value)
        if not row_date:
            continue
        kwh = to_float(ws.cell(row=row_index, column=kwh_col).value)
        conso_eur = to_float(ws.cell(row=row_index, column=conso_col).value) if conso_col else 0.0
        subscription_eur = to_float(ws.cell(row=row_index, column=subscription_col).value) if subscription_col else 0.0
        total_eur = to_float(ws.cell(row=row_index, column=total_col).value) if total_col else 0.0

        plug: float | None = None
        main: float | None = None
        if prod_plug_col and prod_main_col:
            plug = to_float(ws.cell(row=row_index, column=prod_plug_col).value)
            main = to_float(ws.cell(row=row_index, column=prod_main_col).value)
            production = plug + main
        elif prod_total_col:
            production = to_float(ws.cell(row=row_index, column=prod_total_col).value)
        elif prod_single_col:
            production = to_float(ws.cell(row=row_index, column=prod_single_col).value)
            # Before the roof system existed, the single PV column is safely the P&P system.
            if main_service_date and row_date < main_service_date:
                plug = production
        else:
            production = 0.0

        if not any(abs(v) > 0.000001 for v in (kwh, conso_eur, total_eur, production)):
            continue
        rows.append({
            "date": row_date,
            "year": year,
            "month": row_date.month,
            "day": row_date.day,
            "kwh": kwh,
            "conso_eur": conso_eur,
            "subscription_eur": subscription_eur,
            "total_eur": total_eur,
            "production_kwh": production,
            "production_plug_kwh": plug,
            "production_main_kwh": main,
        })
    return rows


def empty_month(year: int, month: int) -> dict[str, Any]:
    return {
        "year": year, "month": month, "label": MONTH_NAMES[month - 1], "days": 0,
        "consumption_kwh": 0.0, "production_kwh": 0.0,
        "production_plug_kwh": 0.0, "production_main_kwh": 0.0,
        "split_days": 0, "conso_eur": 0.0, "subscription_eur": 0.0, "total_eur": 0.0,
        "energy_price_eur_kwh": 0.0, "theoretical_savings_eur": 0.0,
    }


def aggregate_months(rows: list[dict[str, Any]], year: int) -> list[dict[str, Any]]:
    monthly = {m: empty_month(year, m) for m in range(1, 13)}
    for item in rows:
        m = monthly[item["month"]]
        m["days"] += 1
        m["consumption_kwh"] += item["kwh"]
        m["production_kwh"] += item["production_kwh"]
        if item["production_plug_kwh"] is not None and item["production_main_kwh"] is not None:
            m["production_plug_kwh"] += item["production_plug_kwh"]
            m["production_main_kwh"] += item["production_main_kwh"]
            m["split_days"] += 1
        m["conso_eur"] += item["conso_eur"]
        m["subscription_eur"] += item["subscription_eur"]
        m["total_eur"] += item["total_eur"]
    for m in monthly.values():
        if m["consumption_kwh"] > 0:
            m["energy_price_eur_kwh"] = m["conso_eur"] / m["consumption_kwh"]
            m["theoretical_savings_eur"] = m["production_kwh"] * m["energy_price_eur_kwh"]
        if not (m["days"] and m["split_days"] == m["days"]):
            m["production_plug_kwh"] = None
            m["production_main_kwh"] = None
        for key in ["consumption_kwh", "production_kwh", "conso_eur", "subscription_eur", "total_eur", "energy_price_eur_kwh", "theoretical_savings_eur"]:
            m[key] = safe_round(m[key], 4 if "price" in key else 2)
        for key in ["production_plug_kwh", "production_main_kwh"]:
            if m[key] is not None:
                m[key] = safe_round(m[key], 2)
    return [monthly[m] for m in range(1, 13)]


def status_for_year(year: int, latest_year: int, first_full_year: int, phase_dates: list[date]) -> str:
    if year == latest_year:
        return "En cours"
    if year == first_full_year:
        return "Année complète"
    if any(d.year == year for d in phase_dates):
        return "Mise en service"
    if year < first_full_year:
        return "Transition"
    return "Historique"


def build_consumption_matrix(daily_by_year: dict[int, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    matrix = []
    for year, rows in sorted(daily_by_year.items()):
        months = aggregate_months(rows, year)
        output_months = []
        for m in months:
            output_months.append({
                "month": m["month"], "label": m["label"], "days": m["days"],
                "consumption_kwh": m["consumption_kwh"] if m["days"] else None,
            })
        filled = sum(1 for m in output_months if m["days"])
        total = sum((m["consumption_kwh"] or 0) for m in output_months)
        matrix.append({
            "year": year, "filled_months": filled, "months": output_months,
            "total_kwh": safe_round(total, 2) if filled else None,
            "average_month_kwh": safe_round(total / filled, 2) if filled else None,
        })
    return matrix


def read_graf_matrix(workbook: Any) -> dict[int, float]:
    """Read only totals for a consistency check; never use Graf as dashboard truth."""
    if "Graf" not in workbook.sheetnames:
        return {}
    ws = workbook["Graf"]
    out: dict[int, float] = {}
    for r in range(2, ws.max_row + 1):
        year = to_int(ws.cell(r, 1).value, -1)
        if 2000 <= year <= 2100:
            value = to_optional_float(ws.cell(r, 16).value)
            if value is not None:
                out[year] = value
    return out


def sum_through(rows: list[dict[str, Any]], month: int, day: int, key: str) -> float:
    total = 0.0
    for item in rows:
        if (item["date"].month, item["date"].day) <= (month, day):
            total += to_float(item.get(key))
    return total


def build_dashboard_data(workbook_path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=True)
    export = read_pairs(workbook, "EXPORT_SITE")
    params = read_pairs(workbook, "PARAMETRES_COCKPIT", start_row=3)
    kit_years = scan_kit_years(workbook)

    phase_1_date = as_date(export.get("Produit_1_date_mise_service")) or as_date(params.get("plug_service_date"))
    phase_2_date = as_date(export.get("Produit_2_date_mise_service")) or as_date(params.get("main_service_date"))
    phase_dates = [d for d in [phase_1_date, phase_2_date] if d]
    first_install_date = min(phase_dates) if phase_dates else None
    first_full_year = to_int(params.get("first_full_solar_year"), 0)
    if not first_full_year:
        if phase_dates:
            latest_phase = max(phase_dates)
            first_full_year = latest_phase.year if latest_phase.month == 1 and latest_phase.day == 1 else latest_phase.year + 1
        else:
            first_full_year = FALLBACK_FIRST_FULL_SOLAR_YEAR

    daily_by_year: dict[int, list[dict[str, Any]]] = {}
    for sheet_name in workbook.sheetnames:
        if sheet_name.isdigit() and len(sheet_name) == 4:
            year = int(sheet_name)
            rows = read_daily_sheet(workbook, year, phase_2_date)
            if rows:
                daily_by_year[year] = rows
    if not daily_by_year:
        raise ValueError("Aucune feuille annuelle exploitable n'a été trouvée dans le classeur.")

    latest_date = max(item["date"] for rows in daily_by_year.values() for item in rows)
    latest_year = latest_date.year
    if first_install_date is None:
        first_install_date = min(item["date"] for rows in daily_by_year.values() for item in rows)

    plug_power = to_float(params.get("plug_power_kwc"), to_float(export.get("Produit_1_puissance_kwc")))
    main_power = to_float(params.get("main_power_kwc"), to_float(export.get("Produit_2_puissance_kwc")))

    years: list[dict[str, Any]] = []
    monthly_by_year: dict[int, list[dict[str, Any]]] = {}
    daily_json: list[dict[str, Any]] = []
    cumulative_savings = cumulative_production = cumulative_consumption = 0.0
    cumulative_injection = cumulative_autoconsumption = cumulative_total_eur = 0.0
    cumulative_resale_eur = cumulative_net_cost_eur = 0.0

    for year, rows in sorted(daily_by_year.items()):
        months = aggregate_months(rows, year)
        monthly_by_year[year] = months
        production = sum(item["production_kwh"] for item in rows)
        consumption = sum(item["kwh"] for item in rows)
        conso_eur = sum(item["conso_eur"] for item in rows)
        subscription_eur = sum(item["subscription_eur"] for item in rows)
        total_eur = sum(item["total_eur"] for item in rows)
        annual_energy_price = conso_eur / consumption if consumption else 0.0
        theoretical_savings = sum(m["theoretical_savings_eur"] for m in months)

        split_rows = [item for item in rows if item["production_plug_kwh"] is not None and item["production_main_kwh"] is not None]
        split_complete = len(split_rows) == len(rows) and len(rows) > 0
        plug_prod = sum(item["production_plug_kwh"] for item in rows if item["production_plug_kwh"] is not None) if split_complete else None
        main_prod = sum(item["production_main_kwh"] for item in rows if item["production_main_kwh"] is not None) if split_complete else None
        # 2023 is entirely before the roof phase and is therefore P&P-only.
        if phase_2_date and max(item["date"] for item in rows) < phase_2_date:
            plug_prod, main_prod, split_complete = production, 0.0, True

        kit = kit_years.get(year, {})
        injection = kit.get("injection_kwh", 0.0)
        resale_eur = kit.get("resale_eur", 0.0)
        factured_resale_kwh = kit.get("factured_resale_kwh", 0.0)
        autoconsumption = max(production - injection, 0.0)
        workbook_savings = kit.get("workbook_savings_eur", 0.0)
        computed_savings = autoconsumption * annual_energy_price + resale_eur
        realized_savings = workbook_savings if workbook_savings > 0 else computed_savings
        self_consumption_value = max(realized_savings - resale_eur, 0.0)
        actual_home_consumption = consumption + autoconsumption
        autonomy_percent = autoconsumption / actual_home_consumption * 100 if actual_home_consumption else 0.0
        energy_balance_percent = production / actual_home_consumption * 100 if actual_home_consumption else 0.0
        net_energy_balance = production - actual_home_consumption
        net_cost_eur = total_eur - resale_eur

        cumulative_savings += realized_savings
        cumulative_production += production
        cumulative_consumption += consumption
        cumulative_injection += injection
        cumulative_autoconsumption += autoconsumption
        cumulative_total_eur += total_eur
        cumulative_resale_eur += resale_eur
        cumulative_net_cost_eur += net_cost_eur

        first_row_date, last_row_date = min(i["date"] for i in rows), max(i["date"] for i in rows)
        filled_days = len({i["date"] for i in rows})
        year_days = days_in_year(year)
        expected_days_to_latest = (last_row_date - date(year, 1, 1)).days + 1 if year == latest_year else year_days

        plug_yield = plug_prod / plug_power if split_complete and plug_power > 0 else None
        main_yield = main_prod / main_power if split_complete and main_power > 0 and main_prod is not None else None
        plug_vs_main = plug_yield / main_yield * 100 if plug_yield is not None and main_yield and main_yield > 0 else None

        years.append({
            "year": year,
            "status": status_for_year(year, latest_year, first_full_year, phase_dates),
            "is_first_full_year": year == first_full_year,
            "is_current_year": year == latest_year,
            "first_date": date_to_json(first_row_date), "last_date": date_to_json(last_row_date),
            "filled_days": filled_days, "year_days": year_days,
            "coverage_percent": safe_round(filled_days / year_days * 100, 1),
            "coverage_to_date_percent": safe_round(filled_days / expected_days_to_latest * 100, 1) if expected_days_to_latest else 0,
            "consumption_kwh": safe_round(consumption, 2), "production_kwh": safe_round(production, 2),
            "production_plug_kwh": safe_round(plug_prod, 2) if plug_prod is not None else None,
            "production_main_kwh": safe_round(main_prod, 2) if main_prod is not None else None,
            "production_split_available": split_complete,
            "specific_yield_plug_kwh_kwc": safe_round(plug_yield, 1) if plug_yield is not None else None,
            "specific_yield_main_kwh_kwc": safe_round(main_yield, 1) if main_yield is not None else None,
            "plug_vs_main_yield_percent": safe_round(plug_vs_main, 1) if plug_vs_main is not None else None,
            "injection_kwh": safe_round(injection, 2), "autoconsumption_kwh": safe_round(autoconsumption, 2),
            "autoconsumption_percent": safe_round(autoconsumption / production * 100 if production else 0.0, 1),
            "actual_home_consumption_kwh": safe_round(actual_home_consumption, 2),
            "autonomy_percent": safe_round(autonomy_percent, 1),
            "energy_balance_percent": safe_round(energy_balance_percent, 1),
            "net_energy_balance_kwh": safe_round(net_energy_balance, 2),
            "conso_eur": safe_round(conso_eur, 2), "subscription_eur": safe_round(subscription_eur, 2),
            "total_eur": safe_round(total_eur, 2), "resale_eur": safe_round(resale_eur, 2),
            "net_cost_eur": safe_round(net_cost_eur, 2), "factured_resale_kwh": safe_round(factured_resale_kwh, 2),
            "energy_price_eur_kwh": safe_round(annual_energy_price, 4),
            "theoretical_savings_eur": safe_round(theoretical_savings, 2),
            "self_consumption_value_eur": safe_round(self_consumption_value, 2),
            "realized_savings_eur": safe_round(realized_savings, 2),
            "cumulative_savings_eur": safe_round(cumulative_savings, 2),
        })

        for item in rows:
            daily_json.append({
                "date": date_to_json(item["date"]), "year": year, "month": item["month"],
                "consumption_kwh": safe_round(item["kwh"], 2), "production_kwh": safe_round(item["production_kwh"], 2),
                "production_plug_kwh": safe_round(item["production_plug_kwh"], 2) if item["production_plug_kwh"] is not None else None,
                "production_main_kwh": safe_round(item["production_main_kwh"], 2) if item["production_main_kwh"] is not None else None,
                "total_eur": safe_round(item["total_eur"], 2), "conso_eur": safe_round(item["conso_eur"], 2),
            })

    installation_cost = to_float(export.get("cout_installation_Total"))
    remaining_eur = max(installation_cost - cumulative_savings, 0.0)
    amortization_percent = cumulative_savings / installation_cost * 100 if installation_cost else 0.0
    reference_year = next((y for y in years if y["year"] == first_full_year), None)
    if reference_year is None:
        complete = [y for y in years if y["coverage_percent"] >= 99]
        reference_year = complete[-1] if complete else years[-1]
    current = next(y for y in years if y["year"] == latest_year)

    ref_rows = daily_by_year.get(reference_year["year"], [])
    current_rows = daily_by_year[latest_year]
    ref_prod_ytd = sum_through(ref_rows, latest_date.month, latest_date.day, "production_kwh")
    ref_conso_ytd = sum_through(ref_rows, latest_date.month, latest_date.day, "kwh")
    prod_ratio = current["production_kwh"] / ref_prod_ytd if ref_prod_ytd > 0 else 1.0
    conso_ratio = current["consumption_kwh"] / ref_conso_ytd if ref_conso_ytd > 0 else 1.0
    seasonal_prod = reference_year["production_kwh"] * prod_ratio
    seasonal_conso = reference_year["consumption_kwh"] * conso_ratio
    value_per_prod = current["realized_savings_eur"] / current["production_kwh"] if current["production_kwh"] else 0.0
    seasonal_savings = seasonal_prod * value_per_prod if value_per_prod else reference_year["realized_savings_eur"]

    current_days = max(current["filled_days"], 1)
    current_year_days = current["year_days"]
    linear_prod = current["production_kwh"] / current_days * current_year_days
    linear_conso = current["consumption_kwh"] / current_days * current_year_days
    linear_savings = current["realized_savings_eur"] / current_days * current_year_days

    ref_savings = max(reference_year["realized_savings_eur"], 0.0)
    years_ref = remaining_eur / ref_savings if ref_savings > 0 else math.inf
    years_seasonal = remaining_eur / seasonal_savings if seasonal_savings > 0 else math.inf
    payback_ref = add_year_fraction(latest_date, years_ref)
    payback_seasonal = add_year_fraction(latest_date, years_seasonal)

    phases = [
        {"name": "Plug & Play", "date": date_to_json(phase_1_date), "cost_eur": safe_round(to_float(export.get("Produit_1_Pan_Sol_PlugAndPlay")), 2),
         "power_kwc": safe_round(plug_power, 2), "orientation": str(params.get("plug_orientation") or "--"), "role": str(params.get("plug_role") or "")},
        {"name": "Toiture", "date": date_to_json(phase_2_date), "cost_eur": safe_round(to_float(export.get("Produit_2_Pan_Sol_Maison")), 2),
         "power_kwc": safe_round(main_power, 2), "orientation": str(params.get("main_orientation") or "--"), "role": "Production principale"},
    ]
    weighted_investment = weighted_date(phases)
    duration_ref = (payback_ref - weighted_investment).days / 365.25 if payback_ref and weighted_investment else None
    duration_seasonal = (payback_seasonal - weighted_investment).days / 365.25 if payback_seasonal and weighted_investment else None

    # Forecast path for the payback chart. Current year receives an end-of-year estimate,
    # following years use the reference-year savings as a conservative recurring baseline.
    prior_to_current = cumulative_savings - current["realized_savings_eur"]
    forecast_cumulative = prior_to_current + seasonal_savings
    forecast_path = [{"year": latest_year, "cumulative_savings_eur": safe_round(forecast_cumulative, 2)}]
    y = latest_year + 1
    while forecast_cumulative < installation_cost and y <= latest_year + 20:
        forecast_cumulative += ref_savings
        forecast_path.append({"year": y, "cumulative_savings_eur": safe_round(forecast_cumulative, 2)})
        y += 1

    consumption_matrix = build_consumption_matrix(daily_by_year)
    graf_totals = read_graf_matrix(workbook)
    matrix_checks = []
    for row in consumption_matrix:
        graf_total = graf_totals.get(row["year"])
        if graf_total is None or row["total_kwh"] is None:
            continue
        delta = row["total_kwh"] - graf_total
        matrix_checks.append({"year": row["year"], "daily_total_kwh": row["total_kwh"], "graf_total_kwh": safe_round(graf_total, 2), "delta_kwh": safe_round(delta, 2), "ok": abs(delta) < 0.05})

    pool_power_kw = to_float(params.get("pool_pump_power_w")) / 1000
    pool_low_h = to_float(params.get("pool_summer_runtime_low_h"))
    pool_high_h = to_float(params.get("pool_summer_runtime_high_h"))
    pool_low_kwh = pool_power_kw * pool_low_h if pool_power_kw and pool_low_h else to_float(params.get("pool_summer_daily_kwh_low"))
    pool_high_kwh = pool_power_kw * pool_high_h if pool_power_kw and pool_high_h else to_float(params.get("pool_summer_daily_kwh_high"))
    ref_plug_prod = reference_year.get("production_plug_kwh")
    eq_days_low = ref_plug_prod / pool_high_kwh if ref_plug_prod and pool_high_kwh else None
    eq_days_high = ref_plug_prod / pool_low_kwh if ref_plug_prod and pool_low_kwh else None

    return {
        "meta": {
            "version": "3.0", "source_file": workbook_path.name,
            "generated_at": datetime.now().isoformat(timespec="seconds"), "latest_date": date_to_json(latest_date),
            "latest_year": latest_year, "first_full_solar_year": first_full_year, "first_install_date": date_to_json(first_install_date),
        },
        "installation": {
            "name": str(export.get("nom_installation") or "Installation solaire"),
            "power_kwc": safe_round(to_float(export.get("puissance_kwc_Total")), 2), "panel_count": to_int(export.get("nombre_panneaux")),
            "cost_eur": safe_round(installation_cost, 2), "reference_price_eur_kwh": safe_round(to_float(export.get("prix_reference_kwh")), 4),
            "weighted_investment_date": date_to_json(weighted_investment), "phases": phases,
            "main_shading_note": str(params.get("main_shading_note") or ""),
        },
        "pool": {
            "surface_m2_est": safe_round(to_float(params.get("pool_surface_m2_est")), 2), "water_depth_m": safe_round(to_float(params.get("pool_water_depth_m")), 2),
            "volume_m3_est": safe_round(to_float(params.get("pool_volume_m3_est")), 1), "pump_power_w": to_int(params.get("pool_pump_power_w")),
            "summer_runtime_low_h": safe_round(pool_low_h, 1), "summer_runtime_high_h": safe_round(pool_high_h, 1),
            "summer_daily_kwh_low": safe_round(pool_low_kwh, 2), "summer_daily_kwh_high": safe_round(pool_high_kwh, 2),
            "reference_plug_production_kwh": ref_plug_prod, "equivalent_pump_days_low": safe_round(eq_days_low, 0) if eq_days_low else None,
            "equivalent_pump_days_high": safe_round(eq_days_high, 0) if eq_days_high else None,
            "note": "Équivalent énergétique théorique uniquement : la pompe et les panneaux ne coïncident pas toujours heure par heure.",
        },
        "oa": {
            "tariff_eur_kwh": safe_round(to_float(params.get("edf_oa_tariff_eur_kwh")), 4),
            "invoice_2025_start": date_to_json(as_date(params.get("edf_oa_2025_invoice_start"))),
            "invoice_2025_end": date_to_json(as_date(params.get("edf_oa_2025_invoice_end"))),
            "invoice_2025_kwh": safe_round(to_float(params.get("edf_oa_2025_invoice_kwh")), 0),
            "invoice_2025_eur": safe_round(to_float(params.get("edf_oa_2025_invoice_eur")), 2),
            "calendar_note": str(params.get("note_oa_calendar") or ""),
        },
        "totals": {
            "production_kwh": safe_round(cumulative_production, 2), "consumption_kwh": safe_round(cumulative_consumption, 2),
            "injection_kwh": safe_round(cumulative_injection, 2), "autoconsumption_kwh": safe_round(cumulative_autoconsumption, 2),
            "autoconsumption_percent": safe_round(cumulative_autoconsumption / cumulative_production * 100 if cumulative_production else 0.0, 1),
            "electricity_cost_eur": safe_round(cumulative_total_eur, 2), "resale_eur": safe_round(cumulative_resale_eur, 2),
            "net_cost_eur": safe_round(cumulative_net_cost_eur, 2), "realized_savings_eur": safe_round(cumulative_savings, 2),
            "remaining_payback_eur": safe_round(remaining_eur, 2), "amortization_percent": safe_round(amortization_percent, 1),
        },
        "projections": {
            "reference_year": reference_year["year"], "reference_savings_eur": safe_round(ref_savings, 2),
            "reference_production_kwh": reference_year["production_kwh"],
            "current_year": latest_year, "current_days": current_days,
            "reference_ytd_production_kwh": safe_round(ref_prod_ytd, 2), "reference_ytd_consumption_kwh": safe_round(ref_conso_ytd, 2),
            "production_ytd_ratio_vs_reference": safe_round(prod_ratio * 100, 1), "consumption_ytd_ratio_vs_reference": safe_round(conso_ratio * 100, 1),
            "current_seasonal_production_kwh": safe_round(seasonal_prod, 2), "current_seasonal_consumption_kwh": safe_round(seasonal_conso, 2),
            "current_seasonal_savings_eur": safe_round(seasonal_savings, 2),
            "current_linear_production_kwh": safe_round(linear_prod, 2), "current_linear_consumption_kwh": safe_round(linear_conso, 2),
            "current_linear_savings_eur": safe_round(linear_savings, 2),
            # Legacy aliases retained for compatibility with old custom code.
            "current_annualized_production_kwh": safe_round(seasonal_prod, 2), "current_annualized_consumption_kwh": safe_round(seasonal_conso, 2),
            "current_annualized_savings_eur": safe_round(seasonal_savings, 2),
            "years_remaining_reference": None if math.isinf(years_ref) else safe_round(years_ref, 2),
            "years_remaining_seasonal": None if math.isinf(years_seasonal) else safe_round(years_seasonal, 2),
            "payback_reference_date": date_to_json(payback_ref), "payback_seasonal_date": date_to_json(payback_seasonal),
            "weighted_investment_date": date_to_json(weighted_investment),
            "total_payback_duration_reference_years": safe_round(duration_ref, 2) if duration_ref is not None else None,
            "total_payback_duration_seasonal_years": safe_round(duration_seasonal, 2) if duration_seasonal is not None else None,
            "method": f"Projection saisonnalisée par ratio {latest_year} vs {reference_year['year']} à date équivalente; valeur €/kWh produite conservée pour l'économie.",
            "forecast_path": forecast_path,
        },
        "quality": {
            "consumption_matrix_source": "Feuilles quotidiennes (source prioritaire)", "matrix_checks": matrix_checks,
            "monthly_injection_is_estimated": True,
            "monthly_injection_note": str(params.get("note_injection_monthly") or "Ventilation mensuelle estimée faute d'injection mensuelle réelle."),
            "weather_note": str(params.get("note_weather") or "Température/pluie contextualisent mais ne mesurent pas l'irradiation solaire."),
            "oa_calendar_note": str(params.get("note_oa_calendar") or ""),
        },
        "years": years,
        "months": {str(year): months for year, months in monthly_by_year.items()},
        "consumption_matrix": consumption_matrix,
        "daily": daily_json,
    }


def write_dashboard_data(data: dict[str, Any], site_dir: Path) -> Path:
    site_dir.mkdir(parents=True, exist_ok=True)
    output = site_dir / "data.js"
    temp = site_dir / "data.js.tmp"
    payload = json.dumps(data, ensure_ascii=False, indent=2)
    temp.write_text("window.SOLAR_DASHBOARD_DATA = " + payload + ";\n", encoding="utf-8")
    temp.replace(output)
    return output


def main(argv: list[str]) -> int:
    project_dir = Path(__file__).resolve().parents[1]
    workbook_path = Path(argv[0]) if argv else project_dir / "KWH.xlsx"
    if not workbook_path.is_absolute():
        workbook_path = (project_dir / workbook_path).resolve()
    try:
        data = build_dashboard_data(workbook_path)
        output = write_dashboard_data(data, project_dir / "site")
    except Exception as exc:
        print(f"Erreur pendant la génération solaire: {exc}")
        return 1
    print(f"Données solaires mises à jour: {output}")
    print(f"Dernière donnée: {data['meta']['latest_date']}")
    print(f"Amortissement: {data['totals']['amortization_percent']:.1f} %")
    print(f"Projection saisonnalisée {data['meta']['latest_year']}: {data['projections']['current_seasonal_production_kwh']:.0f} kWh")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
